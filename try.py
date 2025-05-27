from pptx_data_processing import (
    process_dap_to_docx,
    process_sop_to_docx,
    process_hmi_to_docx,
    process_scada_to_docx,
)
from docx import Document

from Img_expraction import (
    process_machine_photos,
    process_layout_photos,
    remove_unused_placeholders,
    get_images_from_folder
)

from pdf_doc_extractor import (
    pdf_to_images,
    insert_pdf_images_at_placeholder,
    process_eplan_pdf_to_docx,
    process_alarms_pdf_to_docx
)
from project_details import (
    extract_project_info,
    insert_project_info,
    insert_machine_specifications,
    insert_electrical_specifications
)

doc = Document("template/base_file.docx")
docx_path = "template/base_file.docx"
output_path = "output/manualv2.docx"
pdf_path = "uploads/E-PLAN_Drawing/PM009_EPLAN.pdf"
output_dir = "output/pdf_images"

# process_machine_photos(doc, docx_path) 
# process_machine_photos(doc, docx_path) 
# process_layout_photos(doc, docx_path)
doc, dap_processed = process_dap_to_docx(
    folder_path="uploads/DAP",
    template_path="template/base_file.docx",
)
# process_sop_to_docx(
#     folder_path="uploads/SOP",
#     template_path="template/base_file.docx",
# )
# process_hmi_to_docx(
#     folder_path="uploads/HMI",
#     template_path="template/base_file.docx",
# )
# process_scada_to_docx(
#     folder_path="uploads/SCADA",
#     template_path="template/base_file.docx",
# )

# remove_unused_placeholders(doc, output_path)

# extract_text_from_pdf(pdf_path)
# extract_images_from_pdf(pdf_path, output_dir)
# Example usage
# pdf_to_images(pdf_path, output_dir)

# placeholder= "{{Upload_Electrical_drawing_here}}"
# insert_pdf_images_at_placeholder(docx_path, output_dir, placeholder)

# process_eplan_pdf_to_docx(
#     folder_path = "uploads/E-PLAN_Drawing",
#     template_path = "template/base_file.docx"
# )

# process_alarms_pdf_to_docx(
#     folder_path = "uploads/Alarms",
#     template_path = "template/base_file.docx"
# )



# insert_project_info(
#     docx_path="template/base_file.docx",
#     output_path=docx_path,
# )

# insert_machine_specifications(
#     docx_path="template/base_file.docx",
#     output_path=docx_path,
# )

# insert_electrical_specifications(
#     docx_path="template/base_file.docx",
#     output_path=docx_path,
# )